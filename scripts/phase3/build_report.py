"""Build an editable four-page DOCX from actual fixed-set logs and failure images.

Run with a Python environment containing python-docx (report-only dependency).
"""
from collections import Counter, defaultdict
import csv
from datetime import datetime
from pathlib import Path
import json
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
FOLDER=ROOT/'docs/phase3'
NAMES={'pushup':'Push-up','squat':'Squat','lunge':'Lunge','jumping_jack':'Jumping jack'}


def main():
    log=json.loads((FOLDER/'test_log.json').read_text())
    if 'completed_at_utc' not in log:raise ValueError('Test run incomplete')
    runs=json.loads((FOLDER/'comparison_summary.json').read_text())
    inputs=json.loads((FOLDER/'test_manifest.json').read_text())
    failures=json.loads((FOLDER/'failure_manifest.json').read_text())
    freeze=json.loads((FOLDER/'frozen_configuration.json').read_text())
    scores=json.loads((FOLDER/'reference_metrics.json').read_text())
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.6)
    section.left_margin=section.right_margin=Inches(.65)
    section.header_distance=section.footer_distance=Inches(.25)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.03
    for name,size in [('Title',19),('Heading 1',14),('Heading 2',11)]:
        doc.styles[name].font.name='Arial';doc.styles[name].font.size=Pt(size)
        doc.styles[name].font.color.rgb=RGBColor.from_string('17365D')
    footer=section.footer.paragraphs[0]
    footer.alignment=2
    footer.add_run('POSTURE CHECK | Phase 3 | ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title='POSTURE CHECK — Phase 3 preliminary testing'
    doc.core_properties.author='POSTURE CHECK project team'

    def p(text='',bold=False,size=None):
        paragraph=doc.add_paragraph();run=paragraph.add_run(text);run.bold=bold
        if size:run.font.size=Pt(size)
        return paragraph

    def table(headers,rows,widths=None,size=9):
        t=doc.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
        t.autofit=False
        if widths:
            for col,width in zip(t.columns,widths):col.width=Inches(width)
        for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
        for values in rows:
            cells=t.add_row().cells
            for cell,text in zip(cells,values):cell.text=str(text)
        for row in t.rows:
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(3)
                    for run in paragraph.runs:run.font.size=Pt(size)
        return t

    date=datetime.fromisoformat(log['completed_at_utc']).astimezone(ZoneInfo('Asia/Manila'))
    doc.add_heading('POSTURE CHECK',0)
    p('Phase 3 — Core Implementation & Initial Testing',True,13)
    p(f"Run completed: {date.strftime('%d %B %Y, %I:%M %p')} (Asia/Manila)",size=9)
    doc.add_heading('1. Scope, methods and fixed inputs',1)
    p('The integrated pipeline processes a user-selected exercise: OpenCV decoding → RGB conversion → pretrained MediaPipe Pose Landmarker → visibility/geometry gates → geometric rules → optional temporal post-processing → annotated video and JSON.')
    p('Baseline: independent per-frame rules plus fixed clip aggregation. Main: the same pose model and exercise rules, with feature EMA (0.12 s), phase dwell (0.10 s), feedback dwell (0.15 s), missing-pose reset (0.50 s), and full-cycle repetition estimates. It does not automatically identify exercise type. No custom model was trained.')
    p('The user confirmed the supplied YouTube excerpts were not used to develop or tune this system. Input/model/code hashes were frozen before inference. No parameters were changed. The user subsequently supplied manual references: all four clips correct; counts 4 push-ups, 4 squats, 3 lunges and 10 jumping jacks. These post-run references are not independently expert-verified; source URLs were not supplied.')
    table(['Input','Frames','FPS','Duration','Dimensions'],[[Path(c['file']).name,c['metadata_frames'],f"{c['fps']:.0f}",f"{c['duration_seconds']:.2f} s",f"{c['width']} × {c['height']}"] for c in inputs],size=9)
    p('Total: 1,700 unique input frames; 3,400 frame passes across eight baseline/main runs. Every annotated output was decoded and checked against its input frame count. No run failed.',True)
    doc.add_heading('Conditions and test-set limits',2)
    p('The clips include side/oblique push-ups, side-view squats and lunges, and front-view jumping jacks. Editing/zoom changes, small subjects within padded frames, cropped body parts, and player overlays provide framing/visibility challenges. Controlled poor-lighting, heavy-blur and clutter cases are not adequately represented. Four clips are insufficient for broad generalization claims.')
    p('One-command examples (after environment/model setup):',True)
    p('python run.py --exercise squat --method baseline --input squat.mp4 --output baseline.mp4',size=8)
    p('python run.py --exercise squat --method main --input squat.mp4 --output main.mp4',size=8)

    doc.add_page_break()
    doc.add_heading('2. Preliminary baseline–main comparison',1)
    table(['Exercise','Method','Feature coverage','Processing FPS','State changes¹','Reps²','Clip prediction³'],[
        [NAMES[r['clip_id']],r['method'],f"{100*r['assessable_frames']/r['frames']:.1f}%",f"{r['processing_fps']:.2f}",r['status_changes'],r['repetition_count'] if r['repetition_count'] is not None else 'Not impl.',r['clip_assessment']['prediction']]
        for r in runs],size=8)
    p('¹ All adjacent status changes, including pending and unassessable. These are not classification error counts. Actionable rule_pass ↔ rule_flag flips were zero in every run. Main introduces warm-up/gap states, so total changes can increase.',size=9)
    p('² Compare main estimates with the user-supplied counts below. Baseline has no repetition counter, so its missing value is not zero. Reference counts were supplied after predictions were disclosed.',size=9)
    p('³ Provisional whole-clip rule output, not a verified form label. Jumping-jack clip form classification is intentionally unassessable; open/closed motion can still be tracked.',size=9)
    phases=[]
    for mode in NAMES:
        b=next(r for r in runs if r['clip_id']==mode and r['method']=='baseline')
        m=next(r for r in runs if r['clip_id']==mode and r['method']=='main')
        phases.append(f"{NAMES[mode]} {b['phase_changes']}→{m['phase_changes']}")
    p('Observed phase changes (baseline→main): '+ '; '.join(phases)+'. Fewer phase changes do not prove more accurate phase tracking.')
    runtime=log['runtime']
    p(f"Timing machine: cloud Linux, {runtime['cpu']}, {runtime['logical_cpus']} visible logical CPUs, Python {runtime['python']}, CPU inference. Processing FPS includes inference/features/rules/main temporal updates but excludes startup, decoding, drawing and encoding. Full-run throughput, including those costs, is logged separately. Baseline ran before main for each clip; single-run differences are not evidence of a speed improvement.",size=9)
    doc.add_heading('Reference-based accuracy and counting errors',2)
    p('Both methods agree with 3/4 reference labels (75%, counting the jumping-jack abstention as a miss). Assessed-only agreement is 3/3 (100%), with 75% classification coverage. All references are correct: this cannot measure incorrect-form detection, false accepts or balanced-class performance. Binary macro F1 is not reported.',size=9)
    p('Confusion matrix — identical for baseline and main; rows are user references:',True,size=9)
    table(['Reference','Pred. correct','Pred. incorrect','Abstained'],[
        ['Correct (4 clips)',3,0,1],['Incorrect (0 clips)',0,0,0]
    ],size=8)
    count_rows=[]
    for mode in NAMES:
        row=next(r for r in runs if r['clip_id']==mode and r['method']=='main')
        count_rows.append([NAMES[mode],row['ground_truth_reps'],row['repetition_count'],row['count_absolute_error']])
    table(['Exercise','Reference reps','Main estimate','Absolute error'],count_rows,size=8)
    p(f"Main repetition MAE = {scores['main']['count_mae']:.2f} reps/clip; exact counts on {scores['main']['exact_count_matches']}/4 clips. Baseline counting error is unavailable because counting is not implemented. Laptop FPS remains unmeasured; the timings above are cloud measurements.",True,size=9)

    doc.add_page_break()
    doc.add_heading('3. Failure gallery — ten observed rejection frames',1)
    p('These are ten distinct non-assessable main-run frames from three clips, not ten independent videos or ten proven classification mistakes. They show technical/input failure cases; some abstentions are appropriate. Captions identify observed gate failures. Framing/occlusion explanations are hypotheses unless directly visible.',size=9)
    groups=defaultdict(list)
    for row in failures['examples']:groups[row['category']].append(row)
    order=['Joint confidence','Out-of-frame joints','Size/confidence gates','No pose output']
    descriptions={
        'Joint confidence':'Required joint visibility/presence falls below 0.5.',
        'Out-of-frame joints':'Required predicted landmarks lie outside the image.',
        'Size/confidence gates':'One side fails confidence; the other fails the body-size gate.',
        'No pose output':'No landmarks returned in a cropped/overlay end frame; abstention is appropriate.'}
    for category in order:
        examples=groups.get(category,[])
        if not examples:continue
        p(category+' — '+descriptions[category],True,size=8)
        t=doc.add_table(rows=(len(examples)+2)//3,cols=3);t.autofit=False
        for col in t.columns:col.width=Inches(2.27)
        for index,row in enumerate(examples):
            cell=t.rows[index//3].cells[index%3]
            paragraph=cell.paragraphs[0];paragraph.paragraph_format.space_after=Pt(1)
            paragraph.add_run().add_picture(str(FOLDER/row['file']),width=Inches(2.0))
            caption=cell.add_paragraph(f"F{row['id']:02d} · {NAMES[row['clip']]} · {row['time_seconds']:.2f} s · frame {row['frame_index']}")
            caption.paragraph_format.space_after=Pt(3)
            for run in caption.runs:run.font.size=Pt(7.5)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    p('All images are actual annotated test outputs, not simulations. Full-resolution images, frame indices and diagnosis notes are in failure_manifest.json and failures/. The squat run had no feature rejection; this does not establish correct landmarks or form.',size=8)

    doc.add_page_break()
    doc.add_heading('4. Interpretation, improvement plan and pending work',1)
    p('The main method adds temporal state and movement counts while preserving missing-pose abstention. It does not improve form-label agreement on this set: both methods score 75% overall against the supplied references. All clips are reference-correct, so the result gives no evidence about recognizing incorrect form. All eight actionable-feedback flip counts are zero; total main status changes increase because of warm-up/gap transitions.')
    p('Both methods share the same pose detector and geometry gates, explaining unchanged feature coverage. Main undercounts push-ups by 2, lunges by 3 and jumping jacks by 9; its squat count matches the reference. Short clips, camera edits, missing states and the full-cycle definition are candidate contributors. Count MAE is 3.50 reps/clip. These results were scored post hoc without retuning or rerunning the methods.')
    doc.add_heading('Top three proposed changes (not applied to this test)',2)
    table(['Observed cause','Proposed change','Expected effect / verification'],[
        ['Padding/scale and cropped predicted joints (F01–F04, F09–F10)',
         'Detect active video/person region before relative size checks; separate unsupported/cropped inputs from pose failure.',
         'Fewer framing-related rejections; verify on separate development clips with checked landmarks and count references.'],
        ['Low-confidence required joints / lost pose (F05–F08)',
         'Use view-aware required-joint selection and tracked subject identity; keep honest abstention where reliable angles cannot be measured.',
         'Improve usable pose coverage without accepting equipment/background poses; audit false detections manually.'],
        ['Extra pending/gap transitions and incomplete cycle evidence',
         'Test confidence-aware temporal hysteresis and exercise/variant-specific cycle endpoints on independently counted development clips.',
         'Reduce feedback instability and counting errors; measure count MAE and completion events against manual references.']
    ],size=9)
    doc.add_heading('What is still needed for the hand-in',2)
    p('• Independently review the user-supplied form labels and repetition counts, ideally without viewing predictions. Document the form criteria and partial-cycle convention.\n• Add reference-incorrect clips and a separate fixed hard-case set covering low light, blur, clutter and occlusion. Keep this expanded test set out of tuning.\n• Run the unchanged methods on a laptop and record CPU/RAM/OS/power mode. Copy the code, model, supplied clips and frozen manifest before using the command below.\n• After reviewing results and pending items, commit/tag the submission as cp3-core. No such tag is created by generating this report.',size=9)
    p('Laptop command (environment activated; default manifest/freeze paths):',True,size=9)
    p('python scripts/phase3/run_fixed_tests.py --runtime-label laptop --results docs/phase3/laptop_test_log.json --output-dir outputs/phase3_laptop',size=8)
    doc.add_heading('Reproducibility and source files',2)
    p('test_manifest.json: input hashes and updated references. reference_annotations.json: user reference values, origin and post-run timing; reference_metrics.json: scores and confusion matrices. test_log.json preserves the original eight runs, when labels were still unknown. frozen_configuration.json holds source/model hashes; comparison_summary.json contains merged scores; failure_manifest.json records gallery provenance. Annotated videos/JSON are in outputs/phase3/. Original YouTube channel/URL attribution remains to be supplied.',size=8)
    p('This editable preliminary report distinguishes model predictions, user-supplied references and independently verified measurements. The original clips, inference outputs, method parameters and measured timings are unchanged.',True,size=9)
    output=FOLDER/'Phase3_Testing_Report.docx';doc.save(output)
    with (FOLDER/'comparison_summary.csv').open('w',newline='') as stream:
        writer=csv.writer(stream,lineterminator="\n");writer.writerow(['exercise','method','frames','feature_coverage','processing_fps_cloud','full_run_fps_cloud','state_changes','actionable_flips','phase_changes','predicted_reps','clip_prediction','reference_form','reference_reps'])
        for r in runs:writer.writerow([r['exercise'],r['method'],r['frames'],r['assessable_frames']/r['frames'],r['processing_fps'],r['end_to_end_fps'],r['status_changes'],r['actionable_status_flips'],r['phase_changes'],r['repetition_count'],r['clip_assessment']['prediction'],r['ground_truth_form'],r['ground_truth_reps']])
    print(output)
    print('Inline gallery images:',len(doc.inline_shapes),'explicit page breaks: 3')

if __name__=='__main__':main()
